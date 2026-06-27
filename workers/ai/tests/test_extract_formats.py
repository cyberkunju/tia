"""Universal-ingestion tests: every common timesheet format parses, and the type
is detected from content so a wrong extension/MIME never breaks it."""

from __future__ import annotations

import openpyxl

from tia_ai.extract import extract
from tia_ai.extract.sniff import detect_kind

CLEAN = [
    ["Emp ID", "Full Name", "Working Days", "OT Hours", "Client Code", "Period"],
    ["EMP10001", "Carlos Smith", 22, 5, "CL001", "June 2026"],
    ["EMP10002", "Ahmed Khan", 23, 0, "CL001", "June 2026"],
]


def _xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in CLEAN:
        ws.append(row)
    wb.save(path)
    return path


def _docx(path):
    import docx

    d = docx.Document()
    t = d.add_table(rows=0, cols=len(CLEAN[0]))
    for row in CLEAN:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    d.save(path)
    return path


def _csv(path):
    path.write_text(
        "Emp ID,Full Name,Working Days,OT Hours\n"
        "EMP10001,Carlos Smith,22,5\n"
        "EMP10002,Ahmed Khan,23,0\n",
        encoding="utf-8",
    )
    return path


# ---------------------------------------------------------------- sniffing


def test_detect_xlsx_by_content_even_when_named_png(tmp_path):
    p = _xlsx(tmp_path / "timesheet.png")  # deliberately wrong extension
    assert detect_kind(p, mime="image/png", filename="timesheet.png") == "xlsx"


def test_detect_docx_by_content_even_when_named_txt(tmp_path):
    p = _docx(tmp_path / "sheet.txt")
    assert detect_kind(p, mime="text/plain") == "docx"


def test_detect_pdf_and_image_and_csv(tmp_path):
    pdf = tmp_path / "a.bin"
    pdf.write_bytes(b"%PDF-1.7\n%...")
    assert detect_kind(pdf) == "pdf"
    png = tmp_path / "b.bin"
    png.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 32)
    assert detect_kind(png) == "image"
    assert detect_kind(_csv(tmp_path / "c.csv")) == "csv"


# ---------------------------------------------------------------- extraction


def test_xlsx_named_png_still_extracts(tmp_path):
    # the exact failure from production: a real xlsx delivered as .png
    p = _xlsx(tmp_path / "wa.png")
    r = extract(p, mime="image/png", filename="wa.png")
    assert len(r.rows) == 2
    assert {row.emp_id for row in r.rows} == {"EMP10001", "EMP10002"}


def test_docx_table_extracts(tmp_path):
    r = extract(_docx(tmp_path / "ts.docx"), mime="application/octet-stream")
    assert len(r.rows) == 2
    assert r.rows[0].days_worked == 22


def test_csv_extracts(tmp_path):
    r = extract(_csv(tmp_path / "ts.csv"))
    assert len(r.rows) == 2
    assert r.rows[1].emp_id == "EMP10002"


def test_xlsx_extracts_normally(tmp_path):
    r = extract(_xlsx(tmp_path / "ts.xlsx"))
    assert len(r.rows) == 2
    assert r.client_hint == "CL001" and r.period == "June 2026"


def test_empty_file_is_safe(tmp_path):
    p = tmp_path / "empty.xlsx"
    p.write_bytes(b"")
    assert extract(p).rows == []


def test_garbage_does_not_crash(tmp_path):
    p = tmp_path / "junk.bin"
    p.write_bytes(b"\x00\x01\x02\x03not a real file\xff\xfe")
    assert extract(p).rows == []  # → escalate, never a crash
